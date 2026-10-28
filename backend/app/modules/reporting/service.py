from io import BytesIO
from sqlalchemy import text


async def build_data(db, rtype: str, params: dict) -> dict:
    if rtype == "weekly_kpis":
        rows = (await db.execute(text(
            "SELECT channel, verdict, count(*) c FROM detections WHERE detected_at >= now() - interval '7 days' "
            "GROUP BY 1,2 ORDER BY c DESC"))).all()
        return {"title": "Reporte semanal de KPIs", "text": "Detecciones de los últimos 7 días.",
                "headers": ["Canal", "Veredicto", "Total"], "rows": [[r[0], r[1], r[2]] for r in rows]}
    if rtype == "incident_detail":
        iid = params.get("incident_id")
        inc = (await db.execute(text("SELECT title,status,severity,channel,summary FROM incidents WHERE id=:i"),
                                {"i": iid})).first()
        if not inc:
            raise ValueError("Incidente no encontrado")
        ev = (await db.execute(text("SELECT event_type,created_at FROM incident_events WHERE incident_id=:i "
                                    "ORDER BY created_at"), {"i": iid})).all()
        return {"title": f"Incidente: {inc[0]}",
                "text": f"Estado: {inc[1]} | Severidad: {inc[2]} | Canal: {inc[3]}. "
                        f"{(inc[4] or {}).get('executive_summary', '')}",
                "headers": ["Evento", "Fecha"], "rows": [[e[0], str(e[1])] for e in ev]}
    if rtype == "monthly_management":
        month = params.get("month") or "2026-01"
        rows = (await db.execute(text(
            "SELECT severity,status,count(*) FROM incidents WHERE to_char(created_at,'YYYY-MM')=:m GROUP BY 1,2"),
            {"m": month})).all()
        return {"title": f"Reporte mensual de gestión — {month}", "text": f"Incidentes del mes {month}.",
                "headers": ["Severidad", "Estado", "Total"], "rows": [[r[0], r[1], r[2]] for r in rows]}
    raise ValueError("Tipo de reporte desconocido")


def render(fmt: str, d: dict):
    if fmt == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        buf, st = BytesIO(), getSampleStyleSheet()
        story = [Paragraph(d["title"], st["Title"]), Paragraph(d["text"], st["BodyText"]), Spacer(1, 12)]
        if d["rows"]:
            t = Table([d["headers"]] + d["rows"], repeatRows=1)
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a8a")),
                                   ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                                   ("GRID", (0, 0), (-1, -1), .25, colors.grey)]))
            story.append(t)
        SimpleDocTemplate(buf, pagesize=A4).build(story)
        return buf.getvalue(), "application/pdf"
    if fmt == "docx":
        from docx import Document
        doc = Document()
        doc.add_heading(d["title"], 0)
        doc.add_paragraph(d["text"])
        if d["rows"]:
            t = doc.add_table(rows=1, cols=len(d["headers"]))
            t.style = "Table Grid"
            for i, h in enumerate(d["headers"]):
                t.rows[0].cells[i].text = h
            for r in d["rows"]:
                cells = t.add_row().cells
                for i, v in enumerate(r):
                    cells[i].text = str(v)
        b = BytesIO()
        doc.save(b)
        return b.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if fmt == "xlsx":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Reporte"
        ws.append(d["headers"])
        for r in d["rows"]:
            ws.append(r)
        b = BytesIO()
        wb.save(b)
        return b.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    raise ValueError("Formato no soportado")
